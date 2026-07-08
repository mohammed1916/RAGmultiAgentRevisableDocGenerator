"""Integration tests for the complete document generation pipeline."""

from unittest.mock import Mock, patch
import pytest
import tempfile
import os

from orchestrator import Orchestrator
from models import DocumentRequest
from tools.docx_generator import DOCXGenerator


class TestEndToEndDocumentGeneration:
    """End-to-end integration tests with mocked LLM."""

    @patch("tools.ollama_client.requests.get")
    @patch("tools.ollama_client.requests.post")
    def test_full_document_generation_workflow(self, mock_post, mock_get):
        """Test complete document generation from request to DOCX output."""
        # Mock Ollama connection
        mock_get.return_value.json.return_value = {"models": []}
        mock_get.return_value.status_code = 200

        # Mock LLM responses for each agent

        # 1. Planner response
        planner_response = {
            "document_type": "Technical Specification",
            "assumptions": {
                "audience": "Software engineers",
                "scope": "REST API documentation",
            },
            "tasks": [
                {
                    "id": 1,
                    "description": "Research requirements",
                    "dependencies": [],
                },
                {
                    "id": 2,
                    "description": "Write architecture",
                    "dependencies": [1],
                },
            ],
            "outline": [
                "Executive Summary",
                "Architecture",
                "API Endpoints",
                "Conclusion",
            ],
        }

        # 2. Writer responses (one per section)
        writer_responses = [
            {
                "title": "Executive Summary",
                "content": "This document specifies the REST API architecture and design.",
                "heading_level": 1,
            },
            {
                "title": "Architecture",
                "content": "The system uses microservices architecture with load balancing.",
                "heading_level": 1,
            },
            {
                "title": "API Endpoints",
                "content": "All endpoints follow RESTful conventions with proper HTTP status codes.",
                "heading_level": 1,
            },
            {
                "title": "Conclusion",
                "content": "This specification provides complete guidance for implementation.",
                "heading_level": 1,
            },
        ]

        # 3. Reviewer response (first review)
        review_response = {
            "has_issues": False,
            "grammar_issues": [],
            "consistency_issues": [],
            "structure_issues": [],
            "tone_issues": [],
            "corrections": "Document is well-written and consistent.",
        }

        # 4. Scoring response
        score_response = {
            "relevance": 5,
            "completeness": 4,
            "coherence": 5,
            "structure": 5,
            "overall": 4,
        }

        # Configure mock to return different responses based on call count
        call_count = 0

        def mock_post_side_effect(*args, **kwargs):
            nonlocal call_count
            call_count += 1

            # Alternate between different response types
            if call_count == 1:  # Planner
                return Mock(json=lambda: {"response": str(planner_response)}, status_code=200)
            elif 2 <= call_count <= 5:  # Writer (4 sections)
                section_idx = call_count - 2
                response_data = writer_responses[section_idx] if section_idx < len(writer_responses) else writer_responses[-1]
                return Mock(json=lambda r=response_data: {"response": str(r)}, status_code=200)
            elif call_count == 6:  # Reviewer
                return Mock(json=lambda: {"response": str(review_response)}, status_code=200)
            elif call_count == 7:  # Scorer
                return Mock(json=lambda: {"response": str(score_response)}, status_code=200)
            else:
                return Mock(json=lambda: {"response": "{}"}, status_code=200)

        mock_post.side_effect = mock_post_side_effect

        # Run orchestrator
        orchestrator = Orchestrator()

        request = DocumentRequest(
            request="Create a technical specification for a REST API system"
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            # Temporarily change output directory
            import config
            original_dir = config.config.document_output_dir
            config.config.document_output_dir = tmpdir

            try:
                response = orchestrator.generate_document(request)

                # Verify response structure
                assert response.success is True
                assert response.document_filename is not None
                assert response.execution_plan is not None
                assert response.metrics is not None
                assert response.quality_scores is not None

                # Verify execution plan
                assert response.execution_plan.document_type is not None
                assert len(response.execution_plan.tasks) > 0
                assert len(response.execution_plan.outline) > 0

                # Verify metrics were collected
                assert response.metrics.total_execution_time_ms > 0
                assert response.metrics.planner_latency_ms >= 0
                assert response.metrics.writer_latency_ms >= 0
                assert response.metrics.reviewer_latency_ms >= 0
                assert response.metrics.docx_generation_latency_ms >= 0
                assert response.metrics.num_generated_tasks > 0

                # Verify quality scores
                assert 1 <= response.quality_scores.relevance <= 5
                assert 1 <= response.quality_scores.completeness <= 5
                assert 1 <= response.quality_scores.overall <= 5

                # Verify document file exists
                document_path = os.path.join(tmpdir, response.document_filename)
                assert os.path.exists(document_path)
                assert document_path.endswith(".docx")

            finally:
                config.config.document_output_dir = original_dir

    @patch("tools.ollama_client.requests.get")
    @patch("tools.ollama_client.requests.post")
    def test_document_generation_with_review_issues(self, mock_post, mock_get):
        """Test document generation when reviewer finds issues."""
        # Mock Ollama connection
        mock_get.return_value.json.return_value = {"models": []}
        mock_get.return_value.status_code = 200

        planner_response = {
            "document_type": "Process Documentation",
            "assumptions": {},
            "tasks": [{"id": 1, "description": "Task 1", "dependencies": []}],
            "outline": ["Overview", "Details"],
        }

        writer_response = {
            "title": "Overview",
            "content": "This is an overview.",
            "heading_level": 1,
        }

        # First review finds issues
        review_response_1 = {
            "has_issues": True,
            "grammar_issues": ["Spelling error in line 1"],
            "consistency_issues": ["Inconsistent terminology"],
            "structure_issues": [],
            "tone_issues": [],
            "corrections": "Fix the grammar and terminology issues.",
        }

        score_response = {
            "relevance": 4,
            "completeness": 4,
            "coherence": 4,
            "structure": 4,
            "overall": 4,
        }

        call_count = 0

        def mock_post_side_effect(*args, **kwargs):
            nonlocal call_count
            call_count += 1

            if call_count == 1:  # Planner
                return Mock(json=lambda: {"response": str(planner_response)}, status_code=200)
            elif 2 <= call_count <= 3:  # Writer (2 sections)
                return Mock(json=lambda: {"response": str(writer_response)}, status_code=200)
            elif call_count == 4:  # Reviewer - finds issues
                return Mock(json=lambda: {"response": str(review_response_1)}, status_code=200)
            else:  # Scorer
                return Mock(json=lambda: {"response": str(score_response)}, status_code=200)

        mock_post.side_effect = mock_post_side_effect

        orchestrator = Orchestrator()

        request = DocumentRequest(
            request="Create process documentation"
        )

        with tempfile.TemporaryDirectory() as tmpdir:
            import config
            original_dir = config.config.document_output_dir
            config.config.document_output_dir = tmpdir

            try:
                response = orchestrator.generate_document(request)

                # Should still succeed even with review iterations
                assert response.success is True
                assert response.metrics.review_iterations >= 1

            finally:
                config.config.document_output_dir = original_dir

    def test_docx_generation_quality(self):
        """Test that generated DOCX files have proper structure."""
        gen = DOCXGenerator()
        gen.create_document("Test Report")

        gen.add_heading("Introduction", level=1)
        gen.add_paragraph(
            "This is the introduction section with proper formatting."
        )

        gen.add_heading("Details", level=2)
        gen.add_paragraph("Additional details go here.")
        gen.add_bullet_list(
            [
                "Key point 1",
                "Key point 2",
                "Key point 3",
            ]
        )

        gen.add_page_break()

        gen.add_heading("Conclusion", level=1)
        gen.add_paragraph("Final thoughts and recommendations.")

        with tempfile.TemporaryDirectory() as tmpdir:
            filepath = os.path.join(tmpdir, "test_report.docx")
            saved_path = gen.save(filepath)

            # Verify file exists and has content
            assert os.path.exists(saved_path)
            assert os.path.getsize(saved_path) > 1000  # Should have reasonable size

            # Reload and verify structure
            from docx import Document

            doc = Document(saved_path)
            assert len(doc.paragraphs) > 0
            assert "Test Report" in doc.paragraphs[0].text

    def test_metrics_aggregation(self):
        """Test that metrics are properly aggregated."""
        from tools.metrics import MetricsCollector

        collector = MetricsCollector()
        collector.start_pipeline()

        # Record various metrics
        collector.record_planner_execution(1000.0, 5)
        collector.record_writer_execution(2500.0)
        collector.record_reviewer_execution(800.0, 1)
        collector.record_docx_generation(150.0)

        # Record LLM calls
        collector.record_llm_call(
            "qwen3:8b",
            1000.0,
            prompt_tokens=500,
            completion_tokens=300,
            total_tokens=800,
        )
        collector.record_llm_call(
            "qwen3:8b",
            2500.0,
            prompt_tokens=1000,
            completion_tokens=1200,
            total_tokens=2200,
        )

        metrics = collector.get_pipeline_metrics()

        # Verify metrics
        assert metrics.planner_latency_ms == 1000.0
        assert metrics.writer_latency_ms == 2500.0
        assert metrics.reviewer_latency_ms == 800.0
        assert metrics.num_generated_tasks == 5
        assert metrics.review_iterations == 1
        assert len(metrics.llm_calls) == 2

        # Verify LLM metrics
        summary = collector.get_summary()
        assert summary["total_llm_calls"] == 2
        assert summary["total_prompt_tokens"] == 1500
        assert summary["total_completion_tokens"] == 1500
