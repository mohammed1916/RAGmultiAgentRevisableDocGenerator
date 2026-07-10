"""Microsoft Word document generator."""

import os
from typing import List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from ...exceptions import DOCXGenerationException
from ...logger import setup_logger
from ...models import DocumentSection, DocumentStructure
from .markdown_formatter import MarkdownFormatter

logger = setup_logger(__name__)


class DOCXGenerator:
    """Generate professional Microsoft Word documents."""

    def __init__(self):
        """Initialize the DOCX generator."""
        self.doc = None

    def create_document(self, title: str = None) -> None:
        """Create a new document.

        Args:
            title: Optional document title
        """
        self.doc = Document()
        if title:
            self.add_title(title)
        logger.info("Document created")

    def add_title(self, title: str) -> None:
        """Add document title.

        Args:
            title: Title text
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized. Call create_document() first.")

        heading = self.doc.add_paragraph(title, style="Heading 1")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_format = heading.runs[0].font
        heading_format.size = Pt(24)
        heading_format.bold = True
        heading_format.color.rgb = RGBColor(0, 51, 102)

        logger.debug(f"Added title: {title}")

    def add_heading(self, text: str, level: int = 1) -> None:
        """Add a heading.

        Args:
            text: Heading text
            level: Heading level (1-3)
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        if level < 1 or level > 3:
            raise ValueError("Heading level must be between 1 and 3")

        style = f"Heading {level}"
        self.doc.add_paragraph(text, style=style)
        logger.debug(f"Added heading (level {level}): {text}")

    def add_paragraph(self, text: str, bold: bool = False, italic: bool = False) -> None:
        """Add a paragraph with optional markdown formatting.

        Supports: **bold**, *italic*, ***bold+italic***

        Args:
            text: Paragraph text (may contain markdown formatting)
            bold: Whether entire text should be bold (overrides markdown)
            italic: Whether entire text should be italic (overrides markdown)
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        para = self.doc.add_paragraph()

        # Apply markdown formatting
        segments = MarkdownFormatter.simple_parse(text)
        for segment_text, formatting in segments:
            run = para.add_run(segment_text)
            run.font.bold = bold or formatting.get("bold", False)
            run.font.italic = italic or formatting.get("italic", False)

        logger.debug(f"Added paragraph ({len(text)} chars)")

    def add_markdown_section(self, markdown_text: str) -> None:
        """Add markdown content as properly formatted DOCX elements.

        Parses markdown and creates appropriate DOCX elements:
        - Headers → Heading styles
        - Tables → DOCX tables
        - Lists → DOCX bullet/numbered lists
        - Paragraphs → Normal paragraphs with inline formatting

        Args:
            markdown_text: Markdown-formatted text
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        blocks = MarkdownFormatter.parse_blocks(markdown_text)

        for block in blocks:
            block_type = block.get('type')

            if block_type == 'heading':
                self.add_heading(block['text'], level=min(block['level'], 3))

            elif block_type == 'table':
                rows = block['rows']
                if rows:
                    cols = max(len(row) for row in rows) if rows else 1
                    self.add_table(len(rows), cols, rows)

            elif block_type == 'bullet_list':
                self.add_bullet_list(block['items'])

            elif block_type == 'numbered_list':
                self.add_numbered_list(block['items'])

            elif block_type == 'paragraph':
                self.add_paragraph(block['text'])

            elif block_type == 'code_block':
                # Add code block as monospace paragraph
                para = self.doc.add_paragraph(block.get('code', ''), style='Normal')
                for run in para.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)

        logger.debug(f"Added markdown section with {len(blocks)} blocks")

    def add_bullet_list(self, items: List[str]) -> None:
        """Add a bulleted list with markdown formatting support.

        Supports markdown: **bold**, *italic*, ***bold+italic***

        Args:
            items: List of bullet items (may contain markdown formatting)
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        for item in items:
            para = self.doc.add_paragraph(style="List Bullet")
            # Apply markdown formatting to bullet items
            segments = MarkdownFormatter.simple_parse(item)
            for segment_text, formatting in segments:
                run = para.add_run(segment_text)
                run.font.bold = formatting.get("bold", False)
                run.font.italic = formatting.get("italic", False)

        logger.debug(f"Added bullet list with {len(items)} items")

    def add_numbered_list(self, items: List[str]) -> None:
        """Add a numbered list.

        Args:
            items: List of numbered items
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        for item in items:
            self.doc.add_paragraph(item, style="List Number")

        logger.debug(f"Added numbered list with {len(items)} items")

    def add_table(
        self, rows: int, cols: int, data: List[List[str]] = None
    ) -> None:
        """Add a table.

        Args:
            rows: Number of rows
            cols: Number of columns
            data: Optional table data (list of lists)
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = "Light Grid Accent 1"

        if data:
            for i, row_data in enumerate(data):
                if i < rows:
                    for j, cell_data in enumerate(row_data):
                        if j < cols:
                            table.rows[i].cells[j].text = str(cell_data)

        logger.debug(f"Added table ({rows}x{cols})")

    def add_page_break(self) -> None:
        """Add a page break."""
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        self.doc.add_page_break()
        logger.debug("Added page break")

    def save(self, filepath: str) -> str:
        """Save the document.

        Args:
            filepath: Path where document should be saved

        Returns:
            Absolute path to saved document

        Raises:
            DOCXGenerationException: If document not initialized or save fails
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        try:
            # Create directory if it doesn't exist
            os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

            self.doc.save(filepath)
            abs_path = os.path.abspath(filepath)
            logger.info(f"Document saved: {abs_path}")
            return abs_path
        except Exception as e:
            raise DOCXGenerationException(f"Failed to save document: {str(e)}")

    def from_structure(self, structure: DocumentStructure) -> None:
        """Build document from structured data.

        Args:
            structure: DocumentStructure instance
        """
        self.create_document(structure.title)

        for section in structure.sections:
            if section.heading_level > 0:
                self.add_heading(section.title, section.heading_level)
            # Use markdown section parser to properly handle headers, tables, lists
            self.add_markdown_section(section.content)

        logger.info(f"Document built from structure with {len(structure.sections)} sections")
