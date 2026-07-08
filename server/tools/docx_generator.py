"""Microsoft Word document generator."""

import os
from typing import List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from ..exceptions import DOCXGenerationException
from ..logger import setup_logger
from ..models import DocumentSection, DocumentStructure

logger = setup_logger(__name__)


class DOCXGenerator:
    """Generate professional Microsoft Word documents."""

    def __init__(self):
        """Initialize the DOCX generator."""
        self.doc = None

    def create_document(self, title: str = None) -> None:
        """Create a new document.

        Args:
            title: Optional document title
        """
        self.doc = Document()
        if title:
            self.add_title(title)
        logger.info("Document created")

    def add_title(self, title: str) -> None:
        """Add document title.

        Args:
            title: Title text
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized. Call create_document() first.")

        heading = self.doc.add_paragraph(title, style="Heading 1")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_format = heading.runs[0].font
        heading_format.size = Pt(24)
        heading_format.bold = True
        heading_format.color.rgb = RGBColor(0, 51, 102)

        logger.debug(f"Added title: {title}")

    def add_heading(self, text: str, level: int = 1) -> None:
        """Add a heading.

        Args:
            text: Heading text
            level: Heading level (1-3)
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        if level < 1 or level > 3:
            raise ValueError("Heading level must be between 1 and 3")

        style = f"Heading {level}"
        self.doc.add_paragraph(text, style=style)
        logger.debug(f"Added heading (level {level}): {text}")

    def add_paragraph(self, text: str, bold: bool = False, italic: bool = False) -> None:
        """Add a paragraph.

        Args:
            text: Paragraph text
            bold: Whether text should be bold
            italic: Whether text should be italic
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        para = self.doc.add_paragraph(text)
        if bold or italic:
            for run in para.runs:
                run.font.bold = bold
                run.font.italic = italic

        logger.debug(f"Added paragraph ({len(text)} chars)")

    def add_bullet_list(self, items: List[str]) -> None:
        """Add a bulleted list.

        Args:
            items: List of bullet items
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        for item in items:
            self.doc.add_paragraph(item, style="List Bullet")

        logger.debug(f"Added bullet list with {len(items)} items")

    def add_numbered_list(self, items: List[str]) -> None:
        """Add a numbered list.

        Args:
            items: List of numbered items
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        for item in items:
            self.doc.add_paragraph(item, style="List Number")

        logger.debug(f"Added numbered list with {len(items)} items")

    def add_table(
        self, rows: int, cols: int, data: List[List[str]] = None
    ) -> None:
        """Add a table.

        Args:
            rows: Number of rows
            cols: Number of columns
            data: Optional table data (list of lists)
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        table = self.doc.add_table(rows=rows, cols=cols)
        table.style = "Light Grid Accent 1"

        if data:
            for i, row_data in enumerate(data):
                if i < rows:
                    for j, cell_data in enumerate(row_data):
                        if j < cols:
                            table.rows[i].cells[j].text = str(cell_data)

        logger.debug(f"Added table ({rows}x{cols})")

    def add_page_break(self) -> None:
        """Add a page break."""
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        self.doc.add_page_break()
        logger.debug("Added page break")

    def save(self, filepath: str) -> str:
        """Save the document.

        Args:
            filepath: Path where document should be saved

        Returns:
            Absolute path to saved document

        Raises:
            DOCXGenerationException: If document not initialized or save fails
        """
        if not self.doc:
            raise DOCXGenerationException("Document not initialized.")

        try:
            # Create directory if it doesn't exist
            os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)

            self.doc.save(filepath)
            abs_path = os.path.abspath(filepath)
            logger.info(f"Document saved: {abs_path}")
            return abs_path
        except Exception as e:
            raise DOCXGenerationException(f"Failed to save document: {str(e)}")

    def from_structure(self, structure: DocumentStructure) -> None:
        """Build document from structured data.

        Args:
            structure: DocumentStructure instance
        """
        self.create_document(structure.title)

        for section in structure.sections:
            if section.heading_level > 0:
                self.add_heading(section.title, section.heading_level)
            self.add_paragraph(section.content)

        logger.info(f"Document built from structure with {len(structure.sections)} sections")
